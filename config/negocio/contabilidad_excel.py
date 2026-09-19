"""El libro de ingresos y egresos como .xlsx para mandarle al contador.

Los montos se escriben como números (Decimal), no como texto, para que el
contador pueda sumarlos y filtrarlos.
"""
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .contabilidad import egresos, ingresos, resumen_mensual

MONEDA = '"$"#,##0.00'
FECHA = 'dd/mm/yyyy'
NEGRITA = Font(bold=True)
ESTIMADO = PatternFill('solid', fgColor='FFF3B0')


def _fila(ws, valores):
    """Agrega una fila de datos. El texto viene de afuera (el nombre del cliente
    sale del checkout público): un `=...` se escribiría como fórmula y un
    carácter de control tumba la descarga entera. Se limpia y se fija como texto."""
    ws.append([ILLEGAL_CHARACTERS_RE.sub('', v) if isinstance(v, str) else v
               for v in valores])
    for celda in ws[ws.max_row]:
        if isinstance(celda.value, str):
            celda.data_type = 's'


def _encabezado(ws, columnas):
    ws.append(columnas)
    for celda in ws[ws.max_row]:
        celda.font = NEGRITA


def _fila_total(ws, valores):
    ws.append(valores)
    for celda in ws[ws.max_row]:
        celda.font = NEGRITA


def _formatear(ws, col_fecha, cols_moneda, desde_fila):
    for fila in ws.iter_rows(min_row=desde_fila):
        if col_fecha:
            fila[col_fecha - 1].number_format = FECHA
        for col in cols_moneda:
            fila[col - 1].number_format = MONEDA


def _anchos(ws, anchos):
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def _total(movs):
    return sum((m.monto for m in movs), Decimal('0'))


def generar_xlsx(desde, hasta, titulo):
    ing, egr = ingresos(desde, hasta), egresos(desde, hasta)
    wb = Workbook()

    ws = wb.active
    ws.title = 'Resumen'
    ws.append([f'RYAL — Ingresos y egresos — {titulo}'])
    ws['A1'].font = Font(bold=True, size=13)
    ws.append([])
    _encabezado(ws, ['Mes', 'Ingresos', 'Egresos', 'Diferencia'])
    for f in resumen_mensual(desde, hasta):
        ws.append([f['mes'], f['ingresos'], f['egresos'], f['diferencia']])
    t_ing, t_egr = _total(ing), _total(egr)
    _fila_total(ws, ['Total', t_ing, t_egr, t_ing - t_egr])
    _formatear(ws, None, (2, 3, 4), desde_fila=4)
    ws.freeze_panes = 'A4'
    _anchos(ws, [14, 16, 16, 16])

    ws = wb.create_sheet('Ingresos')
    _encabezado(ws, ['Fecha', 'Origen', 'Referencia', 'Cliente', 'Método de pago',
                     'Monto', 'Pedido cancelado'])
    for m in ing:
        _fila(ws, [m.fecha, m.origen, m.referencia, m.concepto, m.metodo, m.monto,
                   'Sí' if m.cancelado else 'No'])
    _fila_total(ws, ['Total', '', '', '', '', t_ing, ''])
    _formatear(ws, 1, (6,), desde_fila=2)
    ws.freeze_panes = 'A2'
    _anchos(ws, [12, 14, 18, 28, 16, 14, 16])

    ws = wb.create_sheet('Egresos')
    _encabezado(ws, ['Fecha', 'Categoría', 'Referencia', 'Concepto', 'Monto',
                     'Fuente del costo'])
    for m in egr:
        _fila(ws, [m.fecha, m.origen, m.referencia, m.concepto, m.monto, m.fuente_costo])
        if m.estimado:
            for celda in ws[ws.max_row]:
                celda.fill = ESTIMADO
    _fila_total(ws, ['Total', '', '', '', t_egr, ''])
    _formatear(ws, 1, (5,), desde_fila=2)
    ws.freeze_panes = 'A2'
    _anchos(ws, [12, 22, 18, 40, 14, 34])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
