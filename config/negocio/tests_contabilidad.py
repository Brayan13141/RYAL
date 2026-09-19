import datetime
from decimal import Decimal
from zoneinfo import ZoneInfo

from django.test import TestCase

from negocio.models import AjusteCaja, Cliente, Gasto, Pago, Pedido
from orders.models import (
    FUENTE_PRECIO_MENOS_100, FUENTE_REGISTRADO, Order, OrderPayment,
    SupplierOrder, SupplierOrderItem,
)

MX = ZoneInfo('America/Mexico_City')
SEP = (datetime.date(2026, 9, 1), datetime.date(2026, 10, 1))


class Fixtures:
    def _pedido(self, monto='1000', fecha=datetime.date(2026, 9, 10),
                origen=Pedido.TIENDA, estado=Pedido.PAGADO, metodo=Pago.EFECTIVO):
        cli = Cliente.objects.create(nombre='Luis', telefono=f'55{Pedido.objects.count():08d}')
        p = Pedido.objects.create(
            cliente=cli, fecha=fecha, costo_producto=Decimal('600'),
            precio_venta=Decimal(monto), estado=estado, origen=origen)
        Pago.objects.create(pedido=p, fecha=fecha, monto=Decimal(monto), metodo_pago=metodo)
        return p

    def _order(self, codigo, precio='1000', costo='600', qty=1, status='pending'):
        o = Order.objects.create(order_code=codigo, customer_name='Ana',
                                 customer_phone='5512345678', status=status)
        o.items.create(product=None, quantity=qty, price_snapshot=Decimal(precio),
                       cost_snapshot=Decimal(costo) if costo is not None else None,
                       sku_snapshot=f'SKU-{codigo}', name_snapshot='Gorra')
        return o

    def _cobro_web(self, order, monto, fecha=datetime.date(2026, 9, 12)):
        return OrderPayment.objects.create(order=order, fecha=fecha, monto=Decimal(monto),
                                           metodo_pago=OrderPayment.TRANSFERENCIA)

    def _armar(self, order, cuando, estados=('added',)):
        """Arma el pedido al proveedor en `cuando` (datetime aware). Un estado por
        renglón; si hay más estados que renglones, el último renglón repite."""
        so = SupplierOrder.objects.create(order=order, status='done')
        SupplierOrder.objects.filter(pk=so.pk).update(created_at=cuando)  # auto_now_add
        items = list(order.items.order_by('pk'))   # OrderItem no tiene ordering
        for i, estado in enumerate(estados):
            SupplierOrderItem.objects.create(
                supplier_order=so, order_item=items[min(i, len(items) - 1)], status=estado)
        return so


class IngresosTests(Fixtures, TestCase):
    def test_cobros_del_negocio_y_de_la_web_en_el_rango(self):
        from negocio.contabilidad import ingresos
        p = self._pedido('1000', origen=Pedido.WHATSAPP, metodo=Pago.TRANSFERENCIA)
        o = self._order('RYL-1', '800')
        self._cobro_web(o, '500')
        movs = ingresos(*SEP)
        self.assertEqual([(m.origen, m.referencia, m.monto, m.metodo) for m in movs], [
            ('WhatsApp', f'Pedido #{p.pk}', Decimal('1000'), 'Transferencia'),
            ('Web', 'Web RYL-1', Decimal('500'), 'Transferencia'),
        ])
        self.assertEqual(movs[0].concepto, 'Luis')
        self.assertEqual(movs[1].concepto, 'Ana')
        self.assertTrue(all(m.tipo == 'ingreso' for m in movs))

    def test_web_es_bruto_no_solo_la_ganancia(self):
        from negocio.contabilidad import ingresos
        o = self._order('RYL-2', '1000', '600')
        self._cobro_web(o, '1000')
        self.assertEqual(sum(m.monto for m in ingresos(*SEP)), Decimal('1000'))

    def test_fuera_del_rango_no_entra(self):
        from negocio.contabilidad import ingresos
        self._pedido(fecha=datetime.date(2026, 8, 31))
        self._pedido(fecha=datetime.date(2026, 10, 1))
        o = self._order('RYL-3')
        self._cobro_web(o, '300', fecha=datetime.date(2026, 10, 1))
        self.assertEqual(ingresos(*SEP), [])

    def test_cobro_de_pedido_cancelado_entra_marcado(self):
        from negocio.contabilidad import ingresos
        self._pedido(estado=Pedido.CANCELADO)
        o = self._order('RYL-4', status='cancelled')
        self._cobro_web(o, '200')
        movs = ingresos(*SEP)
        self.assertEqual(len(movs), 2)
        self.assertTrue(all(m.cancelado for m in movs))

    def test_pedido_sin_cliente_no_revienta(self):
        from negocio.contabilidad import ingresos
        p = Pedido.objects.create(cliente=None, fecha=datetime.date(2026, 9, 5),
                                  costo_producto=Decimal('0'), precio_venta=Decimal('100'),
                                  origen=Pedido.TIENDA)
        Pago.objects.create(pedido=p, fecha=p.fecha, monto=Decimal('100'))
        self.assertEqual(ingresos(*SEP)[0].concepto, '')


class EgresosTests(Fixtures, TestCase):
    def test_gastos_con_su_categoria(self):
        from negocio.contabilidad import egresos
        g = Gasto.objects.create(fecha=datetime.date(2026, 9, 3), descripcion='Bolsas',
                                 monto=Decimal('150'), categoria=Gasto.OTRO)
        movs = egresos(*SEP)
        self.assertEqual([(m.origen, m.referencia, m.concepto, m.monto) for m in movs],
                         [('Otro', f'Gasto #{g.pk}', 'Bolsas', Decimal('150'))])
        self.assertEqual(movs[0].tipo, 'egreso')

    def test_costo_web_solo_de_renglones_agregados(self):
        from negocio.contabilidad import egresos
        o = self._order('RYL-5', costo='600', qty=2)
        o.items.create(product=None, quantity=1, price_snapshot=Decimal('900'),
                       cost_snapshot=Decimal('500'), sku_snapshot='SKU-X', name_snapshot='Tenis')
        self._armar(o, datetime.datetime(2026, 9, 15, 12, tzinfo=MX),
                    estados=('added', 'variant_not_found'))
        movs = egresos(*SEP)
        self.assertEqual(len(movs), 1)
        m = movs[0]
        self.assertEqual((m.origen, m.referencia, m.concepto, m.monto, m.fuente_costo, m.estimado),
                         ('Costo proveedor web', 'Web RYL-5', 'SKU-RYL-5 — Gorra',
                          Decimal('1200'), FUENTE_REGISTRADO, False))
        self.assertEqual(m.fecha, datetime.date(2026, 9, 15))

    def test_renglon_con_dos_added_cuenta_una_vez(self):
        from negocio.contabilidad import egresos
        o = self._order('RYL-6', costo='600')
        self._armar(o, datetime.datetime(2026, 9, 15, 12, tzinfo=MX), estados=('added', 'added'))
        self.assertEqual(sum(m.monto for m in egresos(*SEP)), Decimal('600'))

    def test_fecha_del_armado_en_hora_de_mexico(self):
        """20:00 del 31 de agosto en México son las 02:00 del 1 de septiembre en UTC."""
        from negocio.contabilidad import egresos
        o = self._order('RYL-7')
        self._armar(o, datetime.datetime(2026, 8, 31, 20, tzinfo=MX))
        self.assertEqual(egresos(*SEP), [])
        movs = egresos(datetime.date(2026, 8, 1), datetime.date(2026, 9, 1))
        self.assertEqual([m.fecha for m in movs], [datetime.date(2026, 8, 31)])

    def test_web_cancelado_y_comprado_si_cuenta(self):
        from negocio.contabilidad import egresos
        o = self._order('RYL-8', costo='600', status='cancelled')
        self._armar(o, datetime.datetime(2026, 9, 15, 12, tzinfo=MX))
        movs = egresos(*SEP)
        self.assertEqual(sum(m.monto for m in movs), Decimal('600'))
        self.assertTrue(movs[0].cancelado)

    def test_costo_estimado_se_marca(self):
        from negocio.contabilidad import egresos
        o = self._order('RYL-9', precio='500', costo=None)
        self._armar(o, datetime.datetime(2026, 9, 15, 12, tzinfo=MX))
        m = egresos(*SEP)[0]
        self.assertEqual((m.monto, m.fuente_costo, m.estimado),
                         (Decimal('400'), FUENTE_PRECIO_MENOS_100, True))

    def test_web_sin_armar_no_genera_egreso(self):
        from negocio.contabilidad import egresos
        o = self._order('RYL-10')
        self._cobro_web(o, '1000')
        self.assertEqual(egresos(*SEP), [])

    def test_costo_producto_del_negocio_y_ajustes_no_aparecen(self):
        """La compra del negocio ya es un Gasto; el arqueo no es un movimiento."""
        from negocio.contabilidad import egresos
        self._pedido('1000')   # costo_producto 600
        AjusteCaja.objects.create(fecha=datetime.date(2026, 9, 5), monto=Decimal('-300'),
                                  saldo_resultante=Decimal('0'), motivo='arqueo')
        self.assertEqual(egresos(*SEP), [])


class ResumenMensualTests(Fixtures, TestCase):
    def test_anual_tiene_12_meses_con_los_ceros(self):
        from negocio.contabilidad import resumen_mensual
        self._pedido('1000', fecha=datetime.date(2026, 3, 10))
        Gasto.objects.create(fecha=datetime.date(2026, 3, 20), descripcion='Renta',
                             monto=Decimal('400'))
        filas = resumen_mensual(datetime.date(2026, 1, 1), datetime.date(2027, 1, 1))
        self.assertEqual(len(filas), 12)
        self.assertEqual(filas[0], {'mes': 'Ene 2026', 'ingresos': Decimal('0'),
                                    'egresos': Decimal('0'), 'diferencia': Decimal('0')})
        self.assertEqual(filas[2], {'mes': 'Mar 2026', 'ingresos': Decimal('1000'),
                                    'egresos': Decimal('400'), 'diferencia': Decimal('600')})
        self.assertEqual(filas[11]['mes'], 'Dic 2026')

    def test_un_mes_es_una_fila(self):
        from negocio.contabilidad import resumen_mensual
        self._pedido('250')
        self.assertEqual(resumen_mensual(*SEP), [
            {'mes': 'Sep 2026', 'ingresos': Decimal('250'),
             'egresos': Decimal('0'), 'diferencia': Decimal('250')}])


class PendientesDeCostoTests(Fixtures, TestCase):
    def test_cuenta_el_web_cobrado_sin_armar(self):
        from negocio.contabilidad import pendientes_de_costo
        sin_armar = self._order('RYL-20')
        self._cobro_web(sin_armar, '500')
        armado = self._order('RYL-21')
        self._cobro_web(armado, '500')
        self._armar(armado, datetime.datetime(2026, 9, 15, 12, tzinfo=MX))
        a_medias = self._order('RYL-22')
        a_medias.items.create(product=None, quantity=1, price_snapshot=Decimal('300'),
                              cost_snapshot=Decimal('100'), sku_snapshot='S', name_snapshot='N')
        self._cobro_web(a_medias, '100')
        self._cobro_web(a_medias, '100')   # dos cobros: cuenta una vez
        self._armar(a_medias, datetime.datetime(2026, 9, 15, 12, tzinfo=MX),
                    estados=('added',))   # solo el primer renglón
        self.assertEqual(pendientes_de_costo(*SEP), 2)

    def test_cancelado_y_fuera_de_rango_no_cuentan(self):
        from negocio.contabilidad import pendientes_de_costo
        cancelado = self._order('RYL-23', status='cancelled')
        self._cobro_web(cancelado, '500')
        viejo = self._order('RYL-24')
        self._cobro_web(viejo, '500', fecha=datetime.date(2026, 8, 30))
        self.assertEqual(pendientes_de_costo(*SEP), 0)


class ExcelTests(Fixtures, TestCase):
    def _libro(self, desde, hasta):
        from io import BytesIO
        from openpyxl import load_workbook
        from negocio.contabilidad_excel import generar_xlsx
        return load_workbook(BytesIO(generar_xlsx(desde, hasta, 'Septiembre 2026')))

    @staticmethod
    def _dec(valor):
        return Decimal(str(valor))

    def test_tres_hojas_y_totales_que_cuadran(self):
        self._pedido('1000')
        o = self._order('RYL-30', precio='800', costo='500')
        self._cobro_web(o, '800')
        self._armar(o, datetime.datetime(2026, 9, 15, 12, tzinfo=MX))
        Gasto.objects.create(fecha=datetime.date(2026, 9, 3), descripcion='Bolsas',
                             monto=Decimal('150.50'))
        wb = self._libro(*SEP)
        self.assertEqual(wb.sheetnames, ['Resumen', 'Ingresos', 'Egresos'])

        ing = wb['Ingresos']
        self.assertEqual([c.value for c in ing[1]], [
            'Fecha', 'Origen', 'Referencia', 'Cliente', 'Método de pago', 'Monto',
            'Pedido cancelado'])
        self.assertEqual(ing.max_row, 4)   # encabezado + 2 cobros + total
        self.assertEqual(ing.cell(4, 1).value, 'Total')
        self.assertEqual(self._dec(ing.cell(4, 6).value), Decimal('1800'))
        self.assertEqual(ing.cell(2, 6).number_format, '"$"#,##0.00')

        egr = wb['Egresos']
        self.assertEqual([c.value for c in egr[1]], [
            'Fecha', 'Categoría', 'Referencia', 'Concepto', 'Monto', 'Fuente del costo'])
        self.assertEqual(self._dec(egr.cell(egr.max_row, 5).value), Decimal('650.50'))

        res = wb['Resumen']
        fila_total = [c.value for c in res[res.max_row]]
        self.assertEqual(fila_total[0], 'Total')
        self.assertEqual([self._dec(v) for v in fila_total[1:]],
                         [Decimal('1800'), Decimal('650.50'), Decimal('1149.50')])

    def test_filas_estimadas_resaltadas(self):
        o = self._order('RYL-31', precio='500', costo=None)
        self._armar(o, datetime.datetime(2026, 9, 15, 12, tzinfo=MX))
        egr = self._libro(*SEP)['Egresos']
        self.assertEqual(egr.cell(2, 6).value, FUENTE_PRECIO_MENOS_100)
        self.assertEqual(egr.cell(2, 1).fill.fgColor.rgb, '00FFF3B0')

    def test_periodo_vacio_es_un_libro_valido_en_cero(self):
        wb = self._libro(*SEP)
        self.assertEqual(self._dec(wb['Ingresos'].cell(2, 6).value), Decimal('0'))
        self.assertEqual(self._dec(wb['Egresos'].cell(2, 5).value), Decimal('0'))
        self.assertEqual(wb['Resumen'].cell(wb['Resumen'].max_row, 1).value, 'Total')


from unittest import mock

from django.contrib.auth.models import User

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


class ContabilidadViewsTests(Fixtures, TestCase):
    URL = '/panel/negocio/contabilidad/'

    def setUp(self):
        User.objects.create_user(username='conta', password='pass', is_staff=True)
        self.client.login(username='conta', password='pass')

    def test_pantalla_del_mes_con_totales(self):
        self._pedido('1000')
        Gasto.objects.create(fecha=datetime.date(2026, 9, 3), descripcion='Bolsas',
                             monto=Decimal('150'))
        res = self.client.get(self.URL, {'mes': '2026-09'})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.context['total_ingresos'], Decimal('1000'))
        self.assertEqual(res.context['total_egresos'], Decimal('150'))
        self.assertEqual(res.context['diferencia'], Decimal('850'))
        self.assertEqual(res.context['periodo_label'], 'Sep 2026')
        self.assertContains(res, 'href="/panel/negocio/contabilidad/excel/?mes=2026-09"')

    def test_anual(self):
        res = self.client.get(self.URL, {'anio': '2026'})
        self.assertEqual(res.context['periodo_label'], '2026')
        self.assertEqual(len(res.context['resumen']), 12)

    def test_parametro_invalido_cae_al_mes_actual(self):
        with mock.patch('negocio.views.timezone.localdate',
                        return_value=datetime.date(2026, 9, 19)):
            res = self.client.get(self.URL, {'mes': 'abc'})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.context['periodo_label'], 'Sep 2026')

    def test_avisos_de_estimados_y_pendientes(self):
        o = self._order('RYL-40', precio='500', costo=None)
        self._armar(o, datetime.datetime(2026, 9, 15, 12, tzinfo=MX))
        sin_armar = self._order('RYL-41')
        self._cobro_web(sin_armar, '500')
        res = self.client.get(self.URL, {'mes': '2026-09'})
        self.assertEqual(res.context['n_estimados'], 1)
        self.assertEqual(res.context['n_pendientes_costo'], 1)

    def test_descarga_excel(self):
        res = self.client.get(self.URL + 'excel/', {'mes': '2026-09'})
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res['Content-Type'], XLSX)
        self.assertEqual(res['Content-Disposition'],
                         'attachment; filename="ryal-contabilidad-2026-09.xlsx"')
        self.assertTrue(res.content.startswith(b'PK'))   # un .xlsx es un zip

    def test_descarga_excel_anual(self):
        res = self.client.get(self.URL + 'excel/', {'anio': '2026'})
        self.assertEqual(res['Content-Disposition'],
                         'attachment; filename="ryal-contabilidad-2026.xlsx"')

    def test_requiere_staff(self):
        self.client.logout()
        self.assertEqual(self.client.get(self.URL).status_code, 302)
        self.assertEqual(self.client.get(self.URL + 'excel/').status_code, 302)


class RevisionFinalTests(Fixtures, TestCase):
    """Hallazgos de la revisión final de rama."""

    def _libro(self):
        from io import BytesIO
        from openpyxl import load_workbook
        from negocio.contabilidad_excel import generar_xlsx
        return load_workbook(BytesIO(generar_xlsx(*SEP, 'Septiembre 2026')))

    def test_nombre_con_formula_queda_como_texto(self):
        """El nombre sale del checkout público: no puede volverse fórmula en el
        Excel que se le manda al contador."""
        o = self._order('RYL-50')
        o.customer_name = '=HYPERLINK("http://x","Ver")'
        o.save()
        self._cobro_web(o, '100')
        celda = self._libro()['Ingresos'].cell(2, 4)
        self.assertEqual(celda.data_type, 's')
        self.assertEqual(celda.value, '=HYPERLINK("http://x","Ver")')

    def test_caracter_de_control_no_rompe_la_descarga(self):
        o = self._order('RYL-51')
        o.customer_name = 'Ana\x0bLópez'
        o.save()
        self._cobro_web(o, '100')
        self.assertEqual(self._libro()['Ingresos'].cell(2, 4).value, 'AnaLópez')

    def test_mismo_dia_ordena_por_numero_de_pedido(self):
        pedidos = [self._pedido('10') for _ in range(11)]
        from negocio.contabilidad import ingresos
        self.assertEqual([m.referencia for m in ingresos(*SEP)],
                         [f'Pedido #{p.pk}' for p in sorted(pedidos, key=lambda p: p.pk)])


class SelectorPeriodoTests(TestCase):
    def setUp(self):
        User.objects.create_user(username='conta2', password='pass', is_staff=True)
        self.client.login(username='conta2', password='pass')

    def test_mes_gana_sobre_anio(self):
        """Desde la vista anual, elegir un mes manda los dos parámetros."""
        res = self.client.get('/panel/negocio/contabilidad/', {'mes': '2026-08', 'anio': '2026'})
        self.assertEqual(res.context['periodo_label'], 'Ago 2026')

    def test_elegir_mes_limpia_el_anio(self):
        res = self.client.get('/panel/negocio/contabilidad/', {'anio': '2026'})
        self.assertContains(res, "this.form.anio.value=''")
